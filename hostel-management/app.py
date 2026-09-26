"""
Hostel Management System
A complete, production-ready Flask application for managing hostel students.
"""
import os
import re
import io
import base64
import secrets
import uuid
from datetime import datetime, date, timezone, timedelta
from functools import wraps

import psycopg2
import psycopg2.extras
import requests

from flask import (
    Flask, render_template, request, redirect, url_for, session,
    flash, send_from_directory, jsonify, abort, g, send_file
)
from werkzeug.utils import secure_filename
from PIL import Image
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
    PageBreak, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass


# --------------------------------------------------------------------------
# CONFIGURATION
# --------------------------------------------------------------------------

BASE_DIR = os.path.abspath(os.path.dirname(__file__))

app = Flask(__name__)

app.config['SECRET_KEY'] = os.environ.get(
    'SECRET_KEY',
    secrets.token_hex(32)
)

app.config['MAX_CONTENT_LENGTH'] = 8 * 1024 * 1024

app.config['UPLOAD_FOLDER'] = os.path.join(
    BASE_DIR,
    'static',
    'uploads'
)


# --------------------------------------------------------------------------
# DATABASE CONFIGURATION
# --------------------------------------------------------------------------

DATABASE_URL = os.environ.get('DATABASE_URL')

if not DATABASE_URL:
    raise RuntimeError(
        'DATABASE_URL is not set. '
        'Add your Supabase connection string to the environment.'
    )


# --------------------------------------------------------------------------
# SUPABASE STORAGE CONFIGURATION
# --------------------------------------------------------------------------

# Remove accidental spaces/newlines from the URL and secret key.
SUPABASE_URL = os.environ.get(
    'SUPABASE_URL',
    ''
).strip().rstrip('/')

SUPABASE_SERVICE_ROLE_KEY = os.environ.get(
    'SUPABASE_SERVICE_ROLE_KEY',
    ''
).strip()

SUPABASE_STORAGE_BUCKET = os.environ.get(
    'SUPABASE_STORAGE_BUCKET',
    'student-photos'
).strip()


if not SUPABASE_URL or not SUPABASE_SERVICE_ROLE_KEY:
    raise RuntimeError(
        'SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY '
        'must be set for photo storage.'
    )


SUPABASE_PUBLIC_PREFIX = (
    f"{SUPABASE_URL}/storage/v1/object/public/"
    f"{SUPABASE_STORAGE_BUCKET}/"
)


# --------------------------------------------------------------------------
# APPLICATION SETTINGS
# --------------------------------------------------------------------------

ADMIN_USERNAME = os.environ.get(
    'ADMIN_USERNAME',
    'admin'
)

ADMIN_PASSWORD = os.environ.get(
    'ADMIN_PASSWORD',
    'admin123'
)

ALLOWED_EXTENSIONS = {'jpg', 'jpeg'}

HOSTELS = [
    'Old Hostel',
    'New Hostel'
]

ROOMS = list(range(1, 11))

SHARINGS = list(range(1, 7))

# Indian Standard Time — every "month end" reset is calculated against
# this timezone, regardless of what timezone the server itself runs in.
IST = timezone(timedelta(hours=5, minutes=30))

# Optional shared-secret for the external cron endpoint
# (/cron/monthly-reset). Set this in your environment and point an
# external scheduler (cron-job.org, UptimeRobot, GitHub Actions, etc.)
# at that URL with ?key=<CRON_SECRET> so the monthly reset fires at
# exactly 12:00 AM IST even if the app is asleep and gets no visitors.
CRON_SECRET = os.environ.get('CRON_SECRET', '').strip()


os.makedirs(
    app.config['UPLOAD_FOLDER'],
    exist_ok=True
)


# --------------------------------------------------------------------------
# DATABASE HELPERS
# --------------------------------------------------------------------------

class DBWrapper:
    """
    Thin wrapper around psycopg2 so the rest of the application
    can use db.execute(...).fetchone().
    """

    def __init__(self, conn):
        self.conn = conn

    def execute(self, query, params=None):
        cur = self.conn.cursor(
            cursor_factory=psycopg2.extras.RealDictCursor
        )
        cur.execute(
            query,
            params or []
        )
        return cur

    def commit(self):
        self.conn.commit()


def get_db():

    if 'db' not in g:

        conn = psycopg2.connect(
            DATABASE_URL
        )

        conn.autocommit = False

        g.db = DBWrapper(conn)

    return g.db


@app.teardown_appcontext
def close_db(exception=None):

    db = g.pop('db', None)

    if db is not None:
        db.conn.close()


def init_db():

    conn = psycopg2.connect(
        DATABASE_URL
    )

    cur = conn.cursor()

    cur.execute('''
        CREATE TABLE IF NOT EXISTS students (
            id SERIAL PRIMARY KEY,
            hostel TEXT NOT NULL
                CHECK(hostel IN ('Old Hostel', 'New Hostel')),
            room_number INTEGER NOT NULL
                CHECK(room_number BETWEEN 1 AND 10),
            sharing INTEGER NOT NULL
                CHECK(sharing BETWEEN 1 AND 6),
            name TEXT NOT NULL,
            contact TEXT NOT NULL,
            total_rent REAL NOT NULL DEFAULT 0,
            amount_paid REAL NOT NULL DEFAULT 0,
            balance REAL NOT NULL DEFAULT 0,
            photo_filename TEXT,
            note TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
    ''')

    cur.execute('''
        CREATE INDEX IF NOT EXISTS idx_hostel_room
        ON students(hostel, room_number)
    ''')

    cur.execute("""
        ALTER TABLE students
        ADD COLUMN IF NOT EXISTS note TEXT NOT NULL DEFAULT ''
    """)

    cur.execute("""
        ALTER TABLE students
        ADD COLUMN IF NOT EXISTS joined_date DATE
    """)

    cur.execute("""
        ALTER TABLE students
        ADD COLUMN IF NOT EXISTS paid_date DATE
    """)

    cur.execute("""
        ALTER TABLE students
        ADD COLUMN IF NOT EXISTS active BOOLEAN NOT NULL DEFAULT TRUE
    """)

    cur.execute('''
        CREATE TABLE IF NOT EXISTS system_meta (
            key TEXT PRIMARY KEY,
            value TEXT
        )
    ''')

    cur.execute('''
        CREATE TABLE IF NOT EXISTS room_notes (
            id SERIAL PRIMARY KEY,
            hostel TEXT NOT NULL
                CHECK(hostel IN ('Old Hostel', 'New Hostel')),
            room_number INTEGER NOT NULL
                CHECK(room_number BETWEEN 1 AND 10),
            note TEXT NOT NULL DEFAULT '',
            updated_at TEXT NOT NULL,
            UNIQUE(hostel, room_number)
        )
    ''')

    conn.commit()

    cur.close()
    conn.close()


# --------------------------------------------------------------------------
# UTILITIES
# --------------------------------------------------------------------------

def allowed_file(filename):

    return (
        '.' in filename
        and filename.rsplit('.', 1)[1].lower()
        in ALLOWED_EXTENSIONS
    )


def validate_mobile(number):

    number = (
        number or ''
    ).strip().replace(' ', '').replace('-', '')

    if number.startswith('+91'):
        number = number[3:]

    elif number.startswith('91') and len(number) == 12:
        number = number[2:]

    return (
        bool(re.fullmatch(r'[6-9]\d{9}', number)),
        number
    )


def ist_now():
    """Current date/time in Indian Standard Time."""
    return datetime.now(IST)


def ist_today():
    """Today's date (IST) as a date object."""
    return ist_now().date()


def ist_month_key(dt=None):
    """'YYYY-MM' key for the given (or current) IST date."""
    dt = dt or ist_now()
    return dt.strftime('%Y-%m')


def parse_date_field(raw_value):
    """
    Parse a 'YYYY-MM-DD' string from an <input type="date"> field.
    Returns (date_or_None, ok). ok is False only when the field was
    non-empty but not a valid date.
    """
    raw_value = (raw_value or '').strip()

    if not raw_value:
        return None, True

    try:
        return datetime.strptime(raw_value, '%Y-%m-%d').date(), True
    except ValueError:
        return None, False


app.jinja_env.globals['today_ist'] = lambda: ist_today().isoformat()


def generate_csrf_token():

    if '_csrf_token' not in session:
        session['_csrf_token'] = secrets.token_hex(16)

    return session['_csrf_token']


app.jinja_env.globals['csrf_token'] = generate_csrf_token
app.jinja_env.globals['HOSTELS'] = HOSTELS
app.jinja_env.globals['ROOMS'] = ROOMS
app.jinja_env.globals['SHARINGS'] = SHARINGS


def validate_csrf(token):

    return (
        token
        and session.get('_csrf_token')
        and secrets.compare_digest(
            token,
            session['_csrf_token']
        )
    )


def login_required(f):

    @wraps(f)
    def decorated(*args, **kwargs):

        if not session.get('admin_logged_in'):

            flash(
                'Please log in to access the admin area.',
                'error'
            )

            return redirect(
                url_for(
                    'admin_login',
                    next=request.path
                )
            )

        return decorated_check(
            f,
            *args,
            **kwargs
        )

    return decorated


def decorated_check(f, *args, **kwargs):
    return f(*args, **kwargs)


# --------------------------------------------------------------------------
# SUPABASE PHOTO UPLOAD
# --------------------------------------------------------------------------

def upload_photo_to_supabase(image_bytes, storage_filename):
    """
    Upload JPEG bytes to Supabase Storage and return the public URL.

    The real Supabase HTTP status and response are printed to the
    VS Code terminal so upload problems can be diagnosed.
    """

    upload_url = (
        f"{SUPABASE_URL}/storage/v1/object/"
        f"{SUPABASE_STORAGE_BUCKET}/"
        f"{storage_filename}"
    )

    headers = {
        'Authorization': f'Bearer {SUPABASE_SERVICE_ROLE_KEY}',
        'apikey': SUPABASE_SERVICE_ROLE_KEY,
        'Content-Type': 'image/jpeg',
        'x-upsert': 'true',
        'Cache-Control': '3600',
    }

    try:

        resp = requests.post(
            upload_url,
            headers=headers,
            data=image_bytes,
            timeout=30
        )

        # --------------------------------------------------------------
        # DEBUG INFORMATION
        # --------------------------------------------------------------

        print()
        print("==========================================")
        print("SUPABASE PHOTO UPLOAD")
        print("==========================================")
        print("URL:", upload_url)
        print("STATUS:", resp.status_code)
        print("RESPONSE:", resp.text[:1000])
        print("==========================================")
        print()

    except requests.RequestException as e:

        print()
        print("==========================================")
        print("SUPABASE CONNECTION ERROR")
        print("==========================================")
        print("ERROR:", repr(e))
        print("==========================================")
        print()

        raise ValueError(
            'Could not reach photo storage. '
            'Please try uploading again.'
        )

    if resp.status_code not in (200, 201):

        raise ValueError(
            f'Photo upload failed. '
            f'Supabase returned HTTP {resp.status_code}. '
            f'Check the VS Code terminal for details.'
        )

    return (
        f"{SUPABASE_PUBLIC_PREFIX}"
        f"{storage_filename}"
    )


# --------------------------------------------------------------------------
# SUPABASE PHOTO DELETE
# --------------------------------------------------------------------------

def delete_photo_from_supabase(photo_value):

    """
    Best-effort delete of a photo from Supabase Storage.
    """

    if not photo_value:
        return

    if not photo_value.startswith(
        SUPABASE_PUBLIC_PREFIX
    ):
        return

    storage_filename = photo_value[
        len(SUPABASE_PUBLIC_PREFIX):
    ]

    delete_url = (
        f"{SUPABASE_URL}/storage/v1/object/"
        f"{SUPABASE_STORAGE_BUCKET}/"
        f"{storage_filename}"
    )

    headers = {
        'Authorization': f'Bearer {SUPABASE_SERVICE_ROLE_KEY}',
        'apikey': SUPABASE_SERVICE_ROLE_KEY,
    }

    try:

        requests.delete(
            delete_url,
            headers=headers,
            timeout=20
        )

    except requests.RequestException:

        pass


# --------------------------------------------------------------------------
# SAVE PHOTO
# --------------------------------------------------------------------------

def save_photo(file_storage):

    """
    Resize/compress uploaded JPG photo and upload it
    to Supabase Storage.
    """

    if (
        not file_storage
        or file_storage.filename == ''
    ):
        return None

    filename = secure_filename(
        file_storage.filename
    )

    if not allowed_file(filename):

        raise ValueError(
            'Only JPG/JPEG images are allowed.'
        )

    try:

        img = Image.open(
            file_storage.stream
        )

        img.verify()

        file_storage.stream.seek(0)

        img = Image.open(
            file_storage.stream
        )

    except Exception:

        raise ValueError(
            'The uploaded file is not a valid image.'
        )

    img = img.convert('RGB')

    img.thumbnail(
        (800, 800),
        Image.LANCZOS
    )

    buf = io.BytesIO()

    img.save(
        buf,
        'JPEG',
        quality=78,
        optimize=True
    )

    image_bytes = buf.getvalue()

    storage_filename = (
        f"{uuid.uuid4().hex}.jpg"
    )

    return upload_photo_to_supabase(
        image_bytes,
        storage_filename
    )


def delete_photo(photo_value):

    delete_photo_from_supabase(
        photo_value
    )


def photo_url(photo_value):

    """
    Resolve stored photo value to displayable URL.
    """

    if not photo_value:
        return None

    if (
        photo_value.startswith('http://')
        or photo_value.startswith('https://')
    ):
        return photo_value

    return url_for(
        'uploaded_file',
        filename=photo_value
    )


# --------------------------------------------------------------------------
# PHOTO PLACEHOLDER
# --------------------------------------------------------------------------

PHOTO_PLACEHOLDER = (
    'data:image/svg+xml;base64,'
    + base64.b64encode(
        b'''
<svg xmlns="http://www.w3.org/2000/svg"
     viewBox="0 0 100 100">
<rect width="100" height="100"
      fill="#E2E8F0"/>
<circle cx="50" cy="38" r="18"
        fill="#94A3B8"/>
<path d="M18 88c0-17.7 14.3-32 32-32s32 14.3 32 32"
      fill="#94A3B8"/>
</svg>
        '''
        .strip()
    )
    .decode('ascii')
)


app.jinja_env.globals['photo_url'] = photo_url
app.jinja_env.globals['PHOTO_PLACEHOLDER'] = PHOTO_PLACEHOLDER


# --------------------------------------------------------------------------
# MONTHLY AUTO-RESET (Indian calendar)
# --------------------------------------------------------------------------
#
# At the start of every new month (00:00 IST, i.e. the moment the
# previous month ends), the following fields are wiped clean for every
# student, ready for that month's payments:
#   - paid_date       ("Date You Paid")
#   - amount_paid      -> 0, balance reset to total_rent
#   - photo_filename   (the receipt photo is deleted from storage too)
#   - note             ("Cash? Paid to whom / payment note")
#
# 'joined_date' is NEVER touched — it is permanent.
#
# This is checked lazily on every request (run_monthly_reset_if_needed)
# so it self-heals even if the server was asleep at midnight. For exact
# timing regardless of traffic, also point an external cron service at
# GET /cron/monthly-reset?key=<CRON_SECRET>.

def run_monthly_reset_if_needed(db):

    current_month = ist_month_key()

    row = db.execute(
        "SELECT value FROM system_meta WHERE key = %s",
        ('last_reset_month',)
    ).fetchone()

    if row and row['value'] == current_month:
        return False

    # Atomically "claim" this month so concurrent requests/workers don't
    # both try to run the reset at the same time.
    cur = db.execute(
        '''
        INSERT INTO system_meta (key, value)
        VALUES ('last_reset_month', %s)
        ON CONFLICT (key) DO UPDATE
            SET value = EXCLUDED.value
            WHERE system_meta.value IS DISTINCT FROM EXCLUDED.value
        RETURNING value
        ''',
        (current_month,)
    )

    claimed = cur.fetchone()

    if not claimed:
        db.commit()
        return False

    students = db.execute(
        "SELECT id, photo_filename FROM students"
    ).fetchall()

    for s in students:

        if s['photo_filename']:

            try:
                delete_photo(s['photo_filename'])
            except Exception:
                pass

    now = datetime.utcnow().isoformat()

    db.execute(
        '''
        UPDATE students
        SET
            paid_date = NULL,
            amount_paid = 0,
            balance = total_rent,
            photo_filename = NULL,
            note = '',
            updated_at = %s
        ''',
        (now,)
    )

    db.commit()

    return True


# --------------------------------------------------------------------------
# ROOM FUNCTIONS
# --------------------------------------------------------------------------

def room_occupancy(
    db,
    hostel,
    room_number,
    exclude_id=None
):

    query = '''
        SELECT
            COUNT(*) FILTER (WHERE active) AS c,
            MAX(sharing) AS s
        FROM students
        WHERE hostel=%s
        AND room_number=%s
    '''

    params = [
        hostel,
        room_number
    ]

    if exclude_id:

        query += '''
            AND id != %s
        '''

        params.append(
            exclude_id
        )

    row = db.execute(
        query,
        params
    ).fetchone()

    return (
        row['c'] or 0,
        row['s']
    )


def check_capacity(
    db,
    hostel,
    room_number,
    sharing,
    exclude_id=None
):

    count, existing_sharing = room_occupancy(
        db,
        hostel,
        room_number,
        exclude_id
    )

    capacity = (
        existing_sharing
        if existing_sharing
        else sharing
    )

    if count >= capacity:

        return (
            False,
            f'{hostel} Room {room_number} is already '
            f'full ({count}/{capacity} occupied). '
            'Please select another room.'
        )

    return True, None


# --------------------------------------------------------------------------
# STATISTICS
# --------------------------------------------------------------------------

def get_stats(db):

    row = db.execute('''
        SELECT
            COUNT(*) AS total,

            SUM(
                CASE
                    WHEN hostel='Old Hostel'
                    THEN 1 ELSE 0
                END
            ) AS old_count,

            SUM(
                CASE
                    WHEN hostel='New Hostel'
                    THEN 1 ELSE 0
                END
            ) AS new_count,

            COALESCE(
                SUM(total_rent),
                0
            ) AS total_rent,

            COALESCE(
                SUM(amount_paid),
                0
            ) AS total_paid,

            COALESCE(
                SUM(balance),
                0
            ) AS total_balance

        FROM students
    ''').fetchone()

    return {
        'total': row['total'] or 0,
        'old_count': row['old_count'] or 0,
        'new_count': row['new_count'] or 0,
        'total_rent': row['total_rent'] or 0,
        'total_paid': row['total_paid'] or 0,
        'total_balance': row['total_balance'] or 0,
    }


# --------------------------------------------------------------------------
# ROOM STRUCTURE
# --------------------------------------------------------------------------

def get_rooms_structure(
    db,
    hostel=None
):

    query = '''
        SELECT *
        FROM students
    '''

    params = []

    if hostel:

        query += '''
            WHERE hostel = %s
        '''

        params.append(hostel)

    query += '''
        ORDER BY LOWER(name)
    '''

    rows = db.execute(
        query,
        params
    ).fetchall()

    structure = {}

    hostels_to_build = (
        [hostel]
        if hostel
        else HOSTELS
    )

    for h in hostels_to_build:

        structure[h] = {
            r: []
            for r in ROOMS
        }

    for row in rows:

        d = dict(row)

        if (
            d['hostel'] in structure
            and d['room_number']
            in structure[d['hostel']]
        ):

            structure[
                d['hostel']
            ][
                d['room_number']
            ].append(d)

    return structure


def row_to_student(row):
    return dict(row)


# --------------------------------------------------------------------------
# INITIALIZE DATABASE
# --------------------------------------------------------------------------

init_db()


@app.before_request
def _auto_monthly_reset():
    """
    Self-healing monthly reset: checked on every request so the reset
    still happens even if no external cron ever fires. Cheap (one
    SELECT) once the month has already been reset.
    """

    if request.endpoint == 'static':
        return

    try:
        run_monthly_reset_if_needed(get_db())
    except Exception:

        try:
            get_db().conn.rollback()
        except Exception:
            pass


@app.route('/cron/monthly-reset')
def cron_monthly_reset():
    """
    Hit this from an external scheduler (cron-job.org, UptimeRobot,
    GitHub Actions cron, etc.) at 00:00 IST daily so the reset fires at
    the exact moment the month ends, even if the app has no visitors
    and/or is asleep (free hosting tiers). Example:

        GET https://your-app-url/cron/monthly-reset?key=YOUR_CRON_SECRET

    Set CRON_SECRET in the environment to enable this endpoint.
    """

    if not CRON_SECRET:
        abort(404)

    if not secrets.compare_digest(
        request.args.get('key', ''),
        CRON_SECRET
    ):
        abort(404)

    db = get_db()
    did_reset = run_monthly_reset_if_needed(db)

    return jsonify({
        'status': 'ok',
        'month': ist_month_key(),
        'reset_ran': did_reset
    })


# --------------------------------------------------------------------------
# PUBLIC ROUTES
# --------------------------------------------------------------------------

@app.route(
    '/',
    methods=['GET', 'POST'],
    endpoint='index'
)
@app.route(
    '/add-student',
    methods=['GET', 'POST'],
    endpoint='add_student'
)
def add_student():

    db = get_db()

    if request.method == 'POST':

        if not validate_csrf(
            request.form.get('csrf_token')
        ):

            flash(
                'Security check failed. Please try again.',
                'error'
            )

            return redirect(
                url_for(request.endpoint)
            )

        hostel = request.form.get(
            'hostel',
            ''
        ).strip()

        room_number = request.form.get(
            'room_number',
            ''
        ).strip()

        sharing = request.form.get(
            'sharing',
            ''
        ).strip()

        name = request.form.get(
            'name',
            ''
        ).strip()

        contact = request.form.get(
            'contact',
            ''
        ).strip()

        total_rent = request.form.get(
            'total_rent',
            ''
        ).strip()

        amount_paid = request.form.get(
            'amount_paid',
            ''
        ).strip()

        note = request.form.get(
            'note',
            ''
        ).strip()[:300]

        joined_date_raw = request.form.get(
            'joined_date',
            ''
        ).strip()

        paid_date_raw = request.form.get(
            'paid_date',
            ''
        ).strip()

        errors = []

        if hostel not in HOSTELS:
            errors.append(
                'Please select a valid hostel.'
            )

        try:

            room_number = int(room_number)

            if room_number not in ROOMS:
                errors.append(
                    'Please select a valid room number.'
                )

        except (ValueError, TypeError):

            errors.append(
                'Please select a valid room number.'
            )

            room_number = None

        try:

            sharing = int(sharing)

            if sharing not in SHARINGS:
                errors.append(
                    'Please select a valid sharing option.'
                )

        except (ValueError, TypeError):

            errors.append(
                'Please select a valid sharing option.'
            )

            sharing = None

        if not name or len(name) < 2:

            errors.append(
                'Please enter a valid student name.'
            )

        mobile_ok, clean_contact = validate_mobile(
            contact
        )

        if not mobile_ok:

            errors.append(
                'Please enter a valid 10-digit Indian mobile number.'
            )

        try:

            total_rent = float(
                total_rent
            )

            if total_rent < 0:
                raise ValueError()

        except (ValueError, TypeError):

            errors.append(
                'Total rent must be a valid non-negative number.'
            )

            total_rent = None

        try:

            amount_paid = float(
                amount_paid
            )

            if amount_paid < 0:
                raise ValueError()

        except (ValueError, TypeError):

            errors.append(
                'Amount paid must be a valid non-negative number.'
            )

            amount_paid = None

        if (
            total_rent is not None
            and amount_paid is not None
            and amount_paid > total_rent
        ):

            errors.append(
                'Amount paid cannot be greater than total rent.'
            )

        joined_date, joined_date_ok = parse_date_field(
            joined_date_raw
        )

        if not joined_date_ok:
            errors.append(
                'Please enter a valid joined date.'
            )

        if joined_date is None and joined_date_ok:
            joined_date = ist_today()

        paid_date, paid_date_ok = parse_date_field(
            paid_date_raw
        )

        if not paid_date_ok:
            errors.append(
                'Please enter a valid payment date.'
            )

        photo_file = request.files.get(
            'photo'
        )

        photo_filename = None

        if (
            not errors
            and room_number
            and sharing
        ):

            ok, msg = check_capacity(
                db,
                hostel,
                room_number,
                sharing
            )

            if not ok:
                errors.append(msg)

        if errors:

            for e in errors:
                flash(e, 'error')

            return render_template(
                'add_student.html',
                form=request.form
            )

        try:

            if (
                photo_file
                and photo_file.filename
            ):

                photo_filename = save_photo(
                    photo_file
                )

        except ValueError as e:

            flash(
                str(e),
                'error'
            )

            return render_template(
                'add_student.html',
                form=request.form
            )

        balance = round(
            total_rent - amount_paid,
            2
        )

        now = datetime.utcnow().isoformat()

        db.execute('''
            INSERT INTO students
            (
                hostel,
                room_number,
                sharing,
                name,
                contact,
                total_rent,
                amount_paid,
                balance,
                photo_filename,
                note,
                joined_date,
                paid_date,
                created_at,
                updated_at
            )
            VALUES
            (
                %s, %s, %s, %s, %s,
                %s, %s, %s, %s, %s,
                %s, %s, %s, %s
            )
        ''', (
            hostel,
            room_number,
            sharing,
            name,
            clean_contact,
            total_rent,
            amount_paid,
            balance,
            photo_filename,
            note,
            joined_date,
            paid_date,
            now,
            now
        ))

        db.commit()

        flash(
            f'{name} was saved successfully to '
            f'{hostel} - Room {room_number}!',
            'success'
        )

        return redirect(
            url_for(
                request.endpoint,
                saved=1,
                hostel=hostel,
                room=room_number,
                name=name
            )
        )

    return render_template(
        'add_student.html',
        form={}
    )


# --------------------------------------------------------------------------
# LEGACY LOCAL UPLOAD ROUTE
# --------------------------------------------------------------------------

@app.route(
    '/static/uploads/<path:filename>'
)
def uploaded_file(filename):

    return send_from_directory(
        app.config['UPLOAD_FOLDER'],
        filename
    )


# --------------------------------------------------------------------------
# ADMIN AUTH
# --------------------------------------------------------------------------

@app.route(
    '/adminlogin',
    methods=['GET', 'POST']
)
def admin_login():

    if session.get('admin_logged_in'):

        return redirect(
            url_for('admin_dashboard')
        )

    if request.method == 'POST':

        if not validate_csrf(
            request.form.get('csrf_token')
        ):

            flash(
                'Security check failed. Please try again.',
                'error'
            )

            return redirect(
                url_for('admin_login')
            )

        username = request.form.get(
            'username',
            ''
        )

        password = request.form.get(
            'password',
            ''
        )

        if (
            secrets.compare_digest(
                username,
                ADMIN_USERNAME
            )
            and
            secrets.compare_digest(
                password,
                ADMIN_PASSWORD
            )
        ):

            session['admin_logged_in'] = True
            session.permanent = True

            flash(
                'Welcome back, admin!',
                'success'
            )

            next_url = request.args.get(
                'next'
            )

            return redirect(
                next_url
                if next_url
                else url_for('admin_dashboard')
            )

        else:

            flash(
                'Invalid username or password.',
                'error'
            )

    return render_template(
        'admin_login.html'
    )


@app.route('/admin/logout')
def admin_logout():

    session.pop(
        'admin_logged_in',
        None
    )

    flash(
        'You have been logged out.',
        'success'
    )

    return redirect(
        url_for('index')
    )


# --------------------------------------------------------------------------
# ADMIN DASHBOARD
# --------------------------------------------------------------------------

@app.route('/admin')
@login_required
def admin_dashboard():

    db = get_db()

    stats = get_stats(db)

    rooms_structure = get_rooms_structure(
        db
    )

    all_students = db.execute(
        '''
        SELECT *
        FROM students
        ORDER BY LOWER(name)
        '''
    ).fetchall()

    all_students = [
        row_to_student(r)
        for r in all_students
    ]

    return render_template(
        'admin.html',
        stats=stats,
        rooms_structure=rooms_structure,
        hostels=HOSTELS,
        rooms=ROOMS,
        all_students=all_students,
    )


# --------------------------------------------------------------------------
# VACANCY
# --------------------------------------------------------------------------

@app.route('/admin/vacancy')
@login_required
def check_vacancy():

    db = get_db()

    students = db.execute(
        '''
        SELECT *
        FROM students
        ORDER BY hostel, room_number, LOWER(name)
        '''
    ).fetchall()

    note_rows = db.execute(
        '''
        SELECT *
        FROM room_notes
        '''
    ).fetchall()

    notes_map = {
        (
            n['hostel'],
            n['room_number']
        ): n['note']
        for n in note_rows
    }

    occupied_rooms = {}

    for row in students:

        s = dict(row)

        key = (
            s['hostel'],
            s['room_number']
        )

        occupied_rooms.setdefault(
            key,
            {
                'sharing': s['sharing'],
                'occupants': []
            }
        )

        occupied_rooms[
            key
        ]['occupants'].append(s)

    hostel_summary = {
        h: {
            'capacity': 0,
            'occupied': 0,
            'vacant': 0
        }
        for h in HOSTELS
    }

    room_cards = []

    for hostel in HOSTELS:

        for room_number in ROOMS:

            key = (
                hostel,
                room_number
            )

            note = notes_map.get(
                key,
                ''
            )

            if key in occupied_rooms:

                info = occupied_rooms[key]

                sharing = info['sharing']

                occupied = sum(
                    1
                    for o in info['occupants']
                    if o.get('active', True)
                )

                vacant = max(
                    sharing - occupied,
                    0
                )

                hostel_summary[
                    hostel
                ]['capacity'] += sharing

                hostel_summary[
                    hostel
                ]['occupied'] += occupied

                hostel_summary[
                    hostel
                ]['vacant'] += vacant

                room_cards.append({
                    'hostel': hostel,
                    'room_number': room_number,
                    'sharing': sharing,
                    'occupied': occupied,
                    'vacant': vacant,
                    'occupants': info['occupants'],
                    'note': note,
                    'has_students': True,
                })

            else:

                room_cards.append({
                    'hostel': hostel,
                    'room_number': room_number,
                    'sharing': None,
                    'occupied': 0,
                    'vacant': None,
                    'occupants': [],
                    'note': note,
                    'has_students': False,
                })

    total_capacity = sum(
        h['capacity']
        for h in hostel_summary.values()
    )

    total_occupied = sum(
        h['occupied']
        for h in hostel_summary.values()
    )

    total_vacant = sum(
        h['vacant']
        for h in hostel_summary.values()
    )

    return render_template(
        'vacancy.html',
        room_cards=room_cards,
        hostel_summary=hostel_summary,
        total_capacity=total_capacity,
        total_occupied=total_occupied,
        total_vacant=total_vacant,
        hostels=HOSTELS,
    )


@app.route(
    '/admin/vacancy/toggle/<int:student_id>',
    methods=['POST']
)
@login_required
def toggle_student_status(student_id):
    """
    Flip a student's ON/OFF status on the vacancy page. OFF frees up
    their bed (it shows as vacant / available) without deleting the
    student's record; ON puts them back as occupying that bed.
    """

    if not validate_csrf(
        request.form.get('csrf_token')
    ):

        flash(
            'Security check failed. Please try again.',
            'error'
        )

        return redirect(
            url_for('check_vacancy')
        )

    db = get_db()

    student = db.execute(
        '''
        SELECT id, name, active
        FROM students
        WHERE id = %s
        ''',
        (student_id,)
    ).fetchone()

    if not student:
        abort(404)

    new_status = not student['active']

    db.execute(
        '''
        UPDATE students
        SET active = %s, updated_at = %s
        WHERE id = %s
        ''',
        (
            new_status,
            datetime.utcnow().isoformat(),
            student_id
        )
    )

    db.commit()

    flash(
        f"{student['name']} is now marked "
        f"{'ON (occupied)' if new_status else 'OFF (vacant)'}.",
        'success'
    )

    return redirect(
        url_for('check_vacancy')
    )


@app.route(
    '/admin/vacancy/note',
    methods=['POST']
)
@login_required
def save_room_note():

    if not validate_csrf(
        request.form.get('csrf_token')
    ):

        flash(
            'Security check failed. Please try again.',
            'error'
        )

        return redirect(
            url_for('check_vacancy')
        )

    hostel = request.form.get(
        'hostel',
        ''
    ).strip()

    room_number = request.form.get(
        'room_number',
        ''
    ).strip()

    note = request.form.get(
        'note',
        ''
    ).strip()[:300]

    if (
        hostel not in HOSTELS
        or not room_number.isdigit()
        or int(room_number) not in ROOMS
    ):

        flash(
            'Invalid room.',
            'error'
        )

        return redirect(
            url_for('check_vacancy')
        )

    room_number = int(room_number)

    db = get_db()

    now = datetime.utcnow().isoformat()

    if note:

        db.execute(
            '''
            INSERT INTO room_notes
            (
                hostel,
                room_number,
                note,
                updated_at
            )
            VALUES
            (%s, %s, %s, %s)

            ON CONFLICT
            (hostel, room_number)

            DO UPDATE SET
                note = EXCLUDED.note,
                updated_at = EXCLUDED.updated_at
            ''',
            (
                hostel,
                room_number,
                note,
                now
            )
        )

        flash(
            f'Note saved for {hostel} - Room {room_number}.',
            'success'
        )

    else:

        db.execute(
            '''
            DELETE FROM room_notes
            WHERE hostel=%s
            AND room_number=%s
            ''',
            (
                hostel,
                room_number
            )
        )

        flash(
            f'Note cleared for {hostel} - Room {room_number}.',
            'success'
        )

    db.commit()

    return redirect(
        url_for('check_vacancy')
    )


# --------------------------------------------------------------------------
# EDIT STUDENT
# --------------------------------------------------------------------------

@app.route(
    '/admin/students/<int:student_id>/edit',
    methods=['GET', 'POST']
)
@login_required
def edit_student(student_id):

    db = get_db()

    student = db.execute(
        '''
        SELECT *
        FROM students
        WHERE id = %s
        ''',
        (student_id,)
    ).fetchone()

    if not student:
        abort(404)

    student = dict(student)

    if request.method == 'POST':

        if not validate_csrf(
            request.form.get('csrf_token')
        ):

            flash(
                'Security check failed. Please try again.',
                'error'
            )

            return redirect(
                url_for(
                    'edit_student',
                    student_id=student_id
                )
            )

        hostel = request.form.get(
            'hostel',
            ''
        ).strip()

        room_number = request.form.get(
            'room_number',
            ''
        ).strip()

        sharing = request.form.get(
            'sharing',
            ''
        ).strip()

        name = request.form.get(
            'name',
            ''
        ).strip()

        contact = request.form.get(
            'contact',
            ''
        ).strip()

        total_rent = request.form.get(
            'total_rent',
            ''
        ).strip()

        amount_paid = request.form.get(
            'amount_paid',
            ''
        ).strip()

        note = request.form.get(
            'note',
            ''
        ).strip()[:300]

        joined_date_raw = request.form.get(
            'joined_date',
            ''
        ).strip()

        paid_date_raw = request.form.get(
            'paid_date',
            ''
        ).strip()

        remove_photo = (
            request.form.get(
                'remove_photo'
            ) == '1'
        )

        errors = []

        if hostel not in HOSTELS:

            errors.append(
                'Please select a valid hostel.'
            )

        try:

            room_number = int(
                room_number
            )

            if room_number not in ROOMS:

                errors.append(
                    'Please select a valid room number.'
                )

        except (ValueError, TypeError):

            errors.append(
                'Please select a valid room number.'
            )

            room_number = None

        try:

            sharing = int(
                sharing
            )

            if sharing not in SHARINGS:

                errors.append(
                    'Please select a valid sharing option.'
                )

        except (ValueError, TypeError):

            errors.append(
                'Please select a valid sharing option.'
            )

            sharing = None

        if not name or len(name) < 2:

            errors.append(
                'Please enter a valid student name.'
            )

        mobile_ok, clean_contact = validate_mobile(
            contact
        )

        if not mobile_ok:

            errors.append(
                'Please enter a valid 10-digit Indian mobile number.'
            )

        try:

            total_rent = float(
                total_rent
            )

            if total_rent < 0:
                raise ValueError()

        except (ValueError, TypeError):

            errors.append(
                'Total rent must be a valid non-negative number.'
            )

            total_rent = None

        try:

            amount_paid = float(
                amount_paid
            )

            if amount_paid < 0:
                raise ValueError()

        except (ValueError, TypeError):

            errors.append(
                'Amount paid must be a valid non-negative number.'
            )

            amount_paid = None

        if (
            total_rent is not None
            and amount_paid is not None
            and amount_paid > total_rent
        ):

            errors.append(
                'Amount paid cannot be greater than total rent.'
            )

        joined_date, joined_date_ok = parse_date_field(
            joined_date_raw
        )

        if not joined_date_ok:
            errors.append(
                'Please enter a valid joined date.'
            )

        if joined_date is None and joined_date_ok:
            # Keep the existing joined date if the field was left blank.
            joined_date = student.get('joined_date') or ist_today()

        paid_date, paid_date_ok = parse_date_field(
            paid_date_raw
        )

        if not paid_date_ok:
            errors.append(
                'Please enter a valid payment date.'
            )

        if (
            not errors
            and room_number
            and sharing
        ):

            ok, msg = check_capacity(
                db,
                hostel,
                room_number,
                sharing,
                exclude_id=student_id
            )

            if not ok:
                errors.append(msg)

        photo_file = request.files.get(
            'photo'
        )

        new_photo_filename = student[
            'photo_filename'
        ]

        if errors:

            for e in errors:
                flash(e, 'error')

            merged = dict(student)
            merged.update(
                request.form
            )

            return render_template(
                'edit_student.html',
                student=merged
            )

        try:

            if (
                photo_file
                and photo_file.filename
            ):

                saved = save_photo(
                    photo_file
                )

                if saved:

                    delete_photo(
                        student['photo_filename']
                    )

                    new_photo_filename = saved

            elif remove_photo:

                delete_photo(
                    student['photo_filename']
                )

                new_photo_filename = None

        except ValueError as e:

            flash(
                str(e),
                'error'
            )

            return render_template(
                'edit_student.html',
                student=student
            )

        balance = round(
            total_rent - amount_paid,
            2
        )

        now = datetime.utcnow().isoformat()

        db.execute(
            '''
            UPDATE students

            SET
                hostel=%s,
                room_number=%s,
                sharing=%s,
                name=%s,
                contact=%s,
                total_rent=%s,
                amount_paid=%s,
                balance=%s,
                photo_filename=%s,
                note=%s,
                joined_date=%s,
                paid_date=%s,
                updated_at=%s

            WHERE id=%s
            ''',
            (
                hostel,
                room_number,
                sharing,
                name,
                clean_contact,
                total_rent,
                amount_paid,
                balance,
                new_photo_filename,
                note,
                joined_date,
                paid_date,
                now,
                student_id
            )
        )

        db.commit()

        flash(
            f'{name} was updated successfully.',
            'success'
        )

        return redirect(
            url_for('admin_dashboard')
        )

    return render_template(
        'edit_student.html',
        student=student
    )


# --------------------------------------------------------------------------
# DELETE STUDENT
# --------------------------------------------------------------------------

@app.route(
    '/admin/students/<int:student_id>/delete',
    methods=['POST']
)
@login_required
def delete_student(student_id):

    if not validate_csrf(
        request.form.get('csrf_token')
    ):

        flash(
            'Security check failed. Please try again.',
            'error'
        )

        return redirect(
            url_for('admin_dashboard')
        )

    db = get_db()

    student = db.execute(
        '''
        SELECT *
        FROM students
        WHERE id = %s
        ''',
        (student_id,)
    ).fetchone()

    if not student:
        abort(404)

    delete_photo(
        student['photo_filename']
    )

    db.execute(
        '''
        DELETE FROM students
        WHERE id = %s
        ''',
        (student_id,)
    )

    db.commit()

    flash(
        f"{student['name']} was deleted.",
        'success'
    )

    return redirect(
        url_for('admin_dashboard')
    )


# --------------------------------------------------------------------------
# EXCEL
# --------------------------------------------------------------------------

EXCEL_COLUMNS = [
    'ID',
    'Hostel',
    'Room Number',
    'Sharing',
    'Name',
    'Contact',
    'Total Rent',
    'Amount Paid',
    'Balance',
    'Photo',
    'Created Date'
]


def _style_header(
    ws,
    row_idx=1
):

    header_fill = PatternFill(
        start_color='1E3A8A',
        end_color='1E3A8A',
        fill_type='solid'
    )

    header_font = Font(
        bold=True,
        color='FFFFFF',
        size=11
    )

    for cell in ws[row_idx]:

        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )


def _autosize(ws):

    for col in ws.columns:

        max_len = 0

        col_letter = (
            col[0].column_letter
        )

        for cell in col:

            try:

                max_len = max(
                    max_len,
                    len(
                        str(cell.value)
                    )
                    if cell.value is not None
                    else 0
                )

            except Exception:
                pass

        ws.column_dimensions[
            col_letter
        ].width = min(
            max(max_len + 4, 10),
            40
        )


def _write_students_sheet(
    ws,
    students
):

    ws.append(
        EXCEL_COLUMNS
    )

    for s in students:

        ws.append([
            s['id'],
            s['hostel'],
            s['room_number'],
            s['sharing'],
            s['name'],
            s['contact'],
            s['total_rent'],
            s['amount_paid'],
            s['balance'],
            s['photo_filename'] or '',
            s['created_at'][:10]
            if s['created_at']
            else ''
        ])

    _style_header(ws)
    _autosize(ws)

    ws.freeze_panes = 'A2'


@app.route('/admin/export-excel')
@login_required
def export_excel():

    db = get_db()

    rows = db.execute(
        '''
        SELECT *
        FROM students
        ORDER BY hostel, room_number, LOWER(name)
        '''
    ).fetchall()

    students = [
        dict(r)
        for r in rows
    ]

    wb = openpyxl.Workbook()

    old_ws = wb.active
    old_ws.title = 'Old Hostel'

    _write_students_sheet(
        old_ws,
        [
            s for s in students
            if s['hostel'] == 'Old Hostel'
        ]
    )

    new_ws = wb.create_sheet(
        'New Hostel'
    )

    _write_students_sheet(
        new_ws,
        [
            s for s in students
            if s['hostel'] == 'New Hostel'
        ]
    )

    summary_ws = wb.create_sheet(
        'Summary'
    )

    stats = get_stats(db)

    summary_ws.append([
        'Metric',
        'Value'
    ])

    summary_ws.append([
        'Total Students',
        stats['total']
    ])

    summary_ws.append([
        'Old Hostel Students',
        stats['old_count']
    ])

    summary_ws.append([
        'New Hostel Students',
        stats['new_count']
    ])

    summary_ws.append([
        'Total Rent (₹)',
        stats['total_rent']
    ])

    summary_ws.append([
        'Total Paid (₹)',
        stats['total_paid']
    ])

    summary_ws.append([
        'Total Balance (₹)',
        stats['total_balance']
    ])

    summary_ws.append([
        'Generated On',
        datetime.now().strftime(
            '%d-%m-%Y %H:%M'
        )
    ])

    _style_header(
        summary_ws
    )

    _autosize(
        summary_ws
    )

    buf = io.BytesIO()

    wb.save(buf)

    buf.seek(0)

    filename = (
        f"hostel_students_"
        f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )

    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        )
    )


# --------------------------------------------------------------------------
# EXCEL IMPORT
# --------------------------------------------------------------------------

@app.route(
    '/admin/import-excel',
    methods=['POST']
)
@login_required
def import_excel():

    if not validate_csrf(
        request.form.get('csrf_token')
    ):

        flash(
            'Security check failed. Please try again.',
            'error'
        )

        return redirect(
            url_for('admin_dashboard')
        )

    file = request.files.get(
        'excel_file'
    )

    if not file or file.filename == '':

        flash(
            'Please choose an Excel file to import.',
            'error'
        )

        return redirect(
            url_for('admin_dashboard')
        )

    if not file.filename.lower().endswith(
        ('.xlsx', '.xlsm')
    ):

        flash(
            'Only .xlsx files are supported for import.',
            'error'
        )

        return redirect(
            url_for('admin_dashboard')
        )

    try:

        wb = openpyxl.load_workbook(
            file,
            data_only=True
        )

        ws = wb.active

    except Exception:

        flash(
            'Could not read the uploaded Excel file. '
            'Please check the format.',
            'error'
        )

        return redirect(
            url_for('admin_dashboard')
        )

    db = get_db()

    header = [
        str(c.value).strip().lower()
        if c.value
        else ''
        for c in ws[1]
    ]

    def col_index(*names):

        for n in names:

            if n in header:
                return header.index(n)

        return None

    idx = {
        'hostel': col_index('hostel'),
        'room': col_index(
            'room number',
            'room_number',
            'room'
        ),
        'sharing': col_index('sharing'),
        'name': col_index('name'),
        'contact': col_index('contact'),
        'rent': col_index(
            'total rent',
            'total_rent',
            'rent'
        ),
        'paid': col_index(
            'amount paid',
            'amount_paid',
            'paid'
        ),
    }

    if any(
        v is None
        for k, v in idx.items()
        if k in (
            'hostel',
            'room',
            'sharing',
            'name',
            'contact',
            'rent',
            'paid'
        )
    ):

        flash(
            'Excel file is missing required columns: '
            'Hostel, Room Number, Sharing, Name, Contact, '
            'Total Rent, Amount Paid.',
            'error'
        )

        return redirect(
            url_for('admin_dashboard')
        )

    imported = 0
    skipped = 0

    now = datetime.utcnow().isoformat()

    for row in ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        try:

            hostel_raw = (
                str(
                    row[idx['hostel']]
                ).strip()
                if row[idx['hostel']]
                else ''
            )

            hostel = next(
                (
                    h for h in HOSTELS
                    if h.lower()
                    == hostel_raw.lower()
                ),
                None
            )

            room_number = int(
                row[idx['room']]
            )

            sharing = int(
                row[idx['sharing']]
            )

            name = (
                str(
                    row[idx['name']]
                ).strip()
                if row[idx['name']]
                else ''
            )

            contact_raw = (
                str(
                    row[idx['contact']]
                ).strip()
                if row[idx['contact']]
                else ''
            )

            total_rent = float(
                row[idx['rent']]
            )

            amount_paid = float(
                row[idx['paid']]
            )

            mobile_ok, clean_contact = (
                validate_mobile(
                    contact_raw
                )
            )

            if (
                not hostel
                or room_number not in ROOMS
                or sharing not in SHARINGS
                or not name
                or len(name) < 2
                or not mobile_ok
                or total_rent < 0
                or amount_paid < 0
                or amount_paid > total_rent
            ):

                skipped += 1
                continue

            ok, _ = check_capacity(
                db,
                hostel,
                room_number,
                sharing
            )

            if not ok:

                skipped += 1
                continue

            balance = round(
                total_rent - amount_paid,
                2
            )

            db.execute(
                '''
                INSERT INTO students
                (
                    hostel,
                    room_number,
                    sharing,
                    name,
                    contact,
                    total_rent,
                    amount_paid,
                    balance,
                    photo_filename,
                    created_at,
                    updated_at
                )
                VALUES
                (
                    %s, %s, %s, %s, %s,
                    %s, %s, %s, NULL, %s, %s
                )
                ''',
                (
                    hostel,
                    room_number,
                    sharing,
                    name,
                    clean_contact,
                    total_rent,
                    amount_paid,
                    balance,
                    now,
                    now
                )
            )

            imported += 1

        except Exception:

            skipped += 1
            continue

    db.commit()

    flash(
        f'Successfully imported: {imported} students. '
        f'Skipped/Invalid: {skipped} rows.',
        'success'
        if imported
        else 'error'
    )

    return redirect(
        url_for('admin_dashboard')
    )


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def _pdf_styles():

    styles = getSampleStyleSheet()

    styles.add(
        ParagraphStyle(
            'MainTitle',
            fontSize=20,
            leading=24,
            alignment=TA_CENTER,
            textColor=colors.HexColor(
                '#1E3A8A'
            ),
            spaceAfter=4,
            fontName='Helvetica-Bold'
        )
    )

    styles.add(
        ParagraphStyle(
            'SubTitle',
            fontSize=10,
            leading=14,
            alignment=TA_CENTER,
            textColor=colors.HexColor(
                '#555555'
            ),
            spaceAfter=12
        )
    )

    styles.add(
        ParagraphStyle(
            'HostelHeading',
            fontSize=16,
            leading=20,
            alignment=TA_LEFT,
            textColor=colors.white,
            fontName='Helvetica-Bold',
            backColor=colors.HexColor(
                '#1E3A8A'
            ),
            spaceAfter=10,
            spaceBefore=14,
            leftIndent=6,
            borderPadding=(
                6, 6, 6, 6
            )
        )
    )

    styles.add(
        ParagraphStyle(
            'RoomHeading',
            fontSize=13,
            leading=16,
            alignment=TA_LEFT,
            textColor=colors.HexColor(
                '#1E3A8A'
            ),
            fontName='Helvetica-Bold',
            spaceBefore=10,
            spaceAfter=4,
            borderWidth=0,
            borderColor=colors.HexColor(
                '#93C5FD'
            ),
            borderPadding=2
        )
    )

    styles.add(
        ParagraphStyle(
            'EmptyRoom',
            fontSize=9.5,
            leading=12,
            alignment=TA_LEFT,
            textColor=colors.HexColor(
                '#888888'
            ),
            fontName='Helvetica-Oblique',
            spaceAfter=8
        )
    )

    styles.add(
        ParagraphStyle(
            'SummaryLabel',
            fontSize=10.5,
            leading=14,
            fontName='Helvetica-Bold',
            textColor=colors.HexColor(
                '#1E3A8A'
            )
        )
    )

    return styles


def _footer(
    canvas_obj,
    doc
):

    canvas_obj.saveState()

    canvas_obj.setFont(
        'Helvetica',
        8
    )

    canvas_obj.setFillColor(
        colors.HexColor('#888888')
    )

    canvas_obj.drawString(
        20 * mm,
        12 * mm,
        f"Generated on "
        f"{datetime.now().strftime('%d-%m-%Y %H:%M')}"
    )

    canvas_obj.drawRightString(
        190 * mm,
        12 * mm,
        f"Page {doc.page}"
    )

    canvas_obj.restoreState()


def _room_table(
    students,
    styles
):

    if not students:

        return Paragraph(
            'No students assigned.',
            styles['EmptyRoom']
        )

    data = [[
        'Name',
        'Contact',
        'Sharing',
        'Total Rent',
        'Paid',
        'Balance'
    ]]

    for s in students:

        data.append([
            s['name'],
            s['contact'],
            f"{s['sharing']} Sharing",
            f"Rs.{s['total_rent']:,.0f}",
            f"Rs.{s['amount_paid']:,.0f}",
            f"Rs.{s['balance']:,.0f}"
        ])

    table = Table(
        data,
        colWidths=[
            38 * mm,
            28 * mm,
            20 * mm,
            24 * mm,
            22 * mm,
            24 * mm
        ],
        repeatRows=1
    )

    style = TableStyle([
        (
            'BACKGROUND',
            (0, 0),
            (-1, 0),
            colors.HexColor('#DBEAFE')
        ),
        (
            'TEXTCOLOR',
            (0, 0),
            (-1, 0),
            colors.HexColor('#1E3A8A')
        ),
        (
            'FONTNAME',
            (0, 0),
            (-1, 0),
            'Helvetica-Bold'
        ),
        (
            'FONTSIZE',
            (0, 0),
            (-1, -1),
            9
        ),
        (
            'GRID',
            (0, 0),
            (-1, -1),
            0.6,
            colors.HexColor('#B0BEC5')
        ),
        (
            'VALIGN',
            (0, 0),
            (-1, -1),
            'MIDDLE'
        ),
        (
            'ALIGN',
            (2, 0),
            (-1, -1),
            'CENTER'
        ),
        (
            'TOPPADDING',
            (0, 0),
            (-1, -1),
            5
        ),
        (
            'BOTTOMPADDING',
            (0, 0),
            (-1, -1),
            5
        ),
        (
            'ROWBACKGROUNDS',
            (0, 1),
            (-1, -1),
            [
                colors.white,
                colors.HexColor('#F8FAFC')
            ]
        ),
    ])

    table.setStyle(style)

    return table


def _build_hostel_section(
    story,
    styles,
    hostel_name,
    rooms_dict
):

    story.append(
        Paragraph(
            hostel_name.upper(),
            styles['HostelHeading']
        )
    )

    for room_no in ROOMS:

        students = rooms_dict.get(
            room_no,
            []
        )

        occ = len(students)

        cap = (
            students[0]['sharing']
            if students
            else None
        )

        cap_text = (
            f" ({occ}/{cap} occupied)"
            if cap
            else
            " (0 occupied)"
        )

        room_block = [
            Paragraph(
                f"ROOM {room_no}{cap_text}",
                styles['RoomHeading']
            ),
            _room_table(
                students,
                styles
            )
        ]

        story.append(
            KeepTogether(
                room_block
            )
        )

        story.append(
            Spacer(1, 4)
        )


def generate_pdf(
    hostel_filter=None
):

    db = get_db()

    buf = io.BytesIO()

    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        leftMargin=16 * mm,
        rightMargin=16 * mm
    )

    styles = _pdf_styles()

    story = []

    if hostel_filter:

        title = (
            f"{hostel_filter.upper()} "
            f"— STUDENT REPORT"
        )

    else:

        title = (
            "HOSTEL MANAGEMENT REPORT"
        )

    story.append(
        Paragraph(
            title,
            styles['MainTitle']
        )
    )

    story.append(
        Paragraph(
            f"Generated Date: "
            f"{datetime.now().strftime('%d-%m-%Y')}",
            styles['SubTitle']
        )
    )

    stats = get_stats(db)

    if not hostel_filter:

        summary_data = [
            [
                'Total Students',
                str(stats['total'])
            ],
            [
                'Total Rent',
                f"Rs.{stats['total_rent']:,.0f}"
            ],
            [
                'Total Paid',
                f"Rs.{stats['total_paid']:,.0f}"
            ],
            [
                'Total Balance',
                f"Rs.{stats['total_balance']:,.0f}"
            ],
        ]

    else:

        rows = db.execute(
            '''
            SELECT
                COUNT(*) c,
                COALESCE(
                    SUM(total_rent),
                    0
                ) r,
                COALESCE(
                    SUM(amount_paid),
                    0
                ) p,
                COALESCE(
                    SUM(balance),
                    0
                ) b
            FROM students
            WHERE hostel=%s
            ''',
            (hostel_filter,)
        ).fetchone()

        summary_data = [
            [
                'Total Students',
                str(rows['c'])
            ],
            [
                'Total Rent',
                f"Rs.{rows['r']:,.0f}"
            ],
            [
                'Total Paid',
                f"Rs.{rows['p']:,.0f}"
            ],
            [
                'Total Balance',
                f"Rs.{rows['b']:,.0f}"
            ],
        ]

    summary_table = Table(
        summary_data,
        colWidths=[
            60 * mm,
            60 * mm
        ]
    )

    summary_table.setStyle(
        TableStyle([
            (
                'BACKGROUND',
                (0, 0),
                (0, -1),
                colors.HexColor('#EFF6FF')
            ),
            (
                'FONTNAME',
                (0, 0),
                (0, -1),
                'Helvetica-Bold'
            ),
            (
                'TEXTCOLOR',
                (0, 0),
                (0, -1),
                colors.HexColor('#1E3A8A')
            ),
            (
                'GRID',
                (0, 0),
                (-1, -1),
                0.6,
                colors.HexColor('#B0BEC5')
            ),
            (
                'FONTSIZE',
                (0, 0),
                (-1, -1),
                10
            ),
            (
                'TOPPADDING',
                (0, 0),
                (-1, -1),
                6
            ),
            (
                'BOTTOMPADDING',
                (0, 0),
                (-1, -1),
                6
            ),
        ])
    )

    story.append(
        summary_table
    )

    story.append(
        Spacer(1, 10)
    )

    if hostel_filter:

        rooms_structure = get_rooms_structure(
            db,
            hostel=hostel_filter
        )

        _build_hostel_section(
            story,
            styles,
            hostel_filter,
            rooms_structure[hostel_filter]
        )

    else:

        rooms_structure = get_rooms_structure(
            db
        )

        for h in HOSTELS:

            _build_hostel_section(
                story,
                styles,
                h,
                rooms_structure[h]
            )

            if h != HOSTELS[-1]:

                story.append(
                    PageBreak()
                )

    doc.build(
        story,
        onFirstPage=_footer,
        onLaterPages=_footer
    )

    buf.seek(0)

    return buf


@app.route('/admin/pdf/old')
@login_required
def pdf_old():

    buf = generate_pdf(
        hostel_filter='Old Hostel'
    )

    return send_file(
        buf,
        as_attachment=True,
        download_name='old_hostel_report.pdf',
        mimetype='application/pdf'
    )


@app.route('/admin/pdf/new')
@login_required
def pdf_new():

    buf = generate_pdf(
        hostel_filter='New Hostel'
    )

    return send_file(
        buf,
        as_attachment=True,
        download_name='new_hostel_report.pdf',
        mimetype='application/pdf'
    )


@app.route('/admin/pdf/complete')
@login_required
def pdf_complete():

    buf = generate_pdf(
        hostel_filter=None
    )

    return send_file(
        buf,
        as_attachment=True,
        download_name='complete_hostel_report.pdf',
        mimetype='application/pdf'
    )


# --------------------------------------------------------------------------
# JSON API
# --------------------------------------------------------------------------

@app.route('/admin/api/students')
@login_required
def api_students():

    db = get_db()

    rows = db.execute(
        '''
        SELECT *
        FROM students
        ORDER BY LOWER(name)
        '''
    ).fetchall()

    return jsonify([
        dict(r)
        for r in rows
    ])


# --------------------------------------------------------------------------
# ERROR HANDLERS
# --------------------------------------------------------------------------

@app.errorhandler(404)
def not_found(e):

    return (
        render_template(
            'error.html',
            code=404,
            message='Page not found.'
        ),
        404
    )


@app.errorhandler(413)
def too_large(e):

    return (
        render_template(
            'error.html',
            code=413,
            message=(
                'The uploaded file is too large. '
                'Maximum size is 8 MB.'
            )
        ),
        413
    )


@app.errorhandler(500)
def server_error(e):

    return (
        render_template(
            'error.html',
            code=500,
            message=(
                'Something went wrong on our end. '
                'Please try again.'
            )
        ),
        500
    )


@app.errorhandler(400)
def bad_request(e):

    return (
        render_template(
            'error.html',
            code=400,
            message='Invalid request.'
        ),
        400
    )


# --------------------------------------------------------------------------
# ENTRYPOINT
# --------------------------------------------------------------------------

if __name__ == '__main__':

    port = int(
        os.environ.get(
            'PORT',
            5000
        )
    )

    debug_mode = (
        os.environ.get(
            'FLASK_DEBUG',
            '0'
        ) == '1'
    )

    app.run(
        host='0.0.0.0',
        port=port,
        debug=debug_mode
    )